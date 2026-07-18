#!/usr/bin/env python3
"""Generate a client-ready NIDS architecture PowerPoint presentation."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

# Brand palette — professional cybersecurity (slate + teal)
NAVY = RGBColor(0x0F, 0x1C, 0x2E)
SLATE = RGBColor(0x1B, 0x2A, 0x41)
TEAL = RGBColor(0x1F, 0xA3, 0xA3)
TEAL_DARK = RGBColor(0x14, 0x7A, 0x7A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
OFF_WHITE = RGBColor(0xF5, 0xF7, 0xFA)
MUTED = RGBColor(0x8A, 0x96, 0xA8)
BODY = RGBColor(0x2D, 0x3A, 0x4A)
LIGHT_LINE = RGBColor(0xD8, 0xDE, 0xE8)
CARD_BG = RGBColor(0xEE, 0xF2, 0xF7)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

OUTPUT = Path(__file__).resolve().parent / "NIDS_Client_Presentation.pptx"
ARCH_IMG = Path(__file__).resolve().parent / "assets" / "nids-architecture-diagram.png"


def set_run(run, text, size=18, bold=False, color=BODY, font="Calibri"):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_bg(slide, color=OFF_WHITE):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    # send to back
    spTree = slide.shapes._spTree
    sp = shape._element
    spTree.remove(sp)
    spTree.insert(2, sp)


def add_top_bar(slide):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Inches(0.12))
    bar.fill.solid()
    bar.fill.fore_color.rgb = TEAL
    bar.line.fill.background()


def add_footer(slide, page, total):
    box = slide.shapes.add_textbox(Inches(0.6), Inches(7.05), Inches(10), Inches(0.3))
    tf = box.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    set_run(run, "Network Intrusion Detection System  |  Confidential — Client Briefing", 10, False, MUTED)
    num = slide.shapes.add_textbox(Inches(11.8), Inches(7.05), Inches(1.2), Inches(0.3))
    tf2 = num.text_frame
    p2 = tf2.paragraphs[0]
    p2.alignment = PP_ALIGN.RIGHT
    run2 = p2.add_run()
    set_run(run2, f"{page} / {total}", 10, False, MUTED)


def add_title_block(slide, title, subtitle=None):
    add_top_bar(slide)
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(0.45), Inches(0.12), Inches(0.55))
    accent.fill.solid()
    accent.fill.fore_color.rgb = TEAL
    accent.line.fill.background()

    box = slide.shapes.add_textbox(Inches(0.9), Inches(0.4), Inches(11.5), Inches(0.55))
    tf = box.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    set_run(run, title, 28, True, NAVY)

    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.9), Inches(0.95), Inches(11.5), Inches(0.4))
        tf2 = sub.text_frame
        p2 = tf2.paragraphs[0]
        run2 = p2.add_run()
        set_run(run2, subtitle, 14, False, MUTED)


def add_card(slide, left, top, width, height, title, body_lines, icon_color=TEAL):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = WHITE
    card.line.color.rgb = LIGHT_LINE
    card.line.width = Pt(1)
    try:
        card.adjustments[0] = 0.08
    except Exception:
        pass

    stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(0.08), height)
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = icon_color
    stripe.line.fill.background()

    tbox = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.2), width - Inches(0.4), Inches(0.4))
    tf = tbox.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    set_run(run, title, 16, True, NAVY)

    bbox = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.65), width - Inches(0.4), height - Inches(0.85))
    tf2 = bbox.text_frame
    tf2.word_wrap = True
    for i, line in enumerate(body_lines):
        p = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
        p.space_after = Pt(4)
        run = p.add_run()
        set_run(run, line, 13, False, BODY)


def blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])  # blank


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    total = 12

    # ----- 1. Title -----
    s = blank_slide(prs)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    bg.line.fill.background()

    # left accent panel
    panel = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.25), SLIDE_H)
    panel.fill.solid()
    panel.fill.fore_color.rgb = TEAL
    panel.line.fill.background()

    eyebrow = s.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11), Inches(0.4))
    p = eyebrow.text_frame.paragraphs[0]
    run = p.add_run()
    set_run(run, "SECURITY ARCHITECTURE BRIEFING", 14, True, TEAL)

    title = s.shapes.add_textbox(Inches(1.0), Inches(2.5), Inches(11), Inches(1.2))
    p = title.text_frame.paragraphs[0]
    run = p.add_run()
    set_run(run, "Network Intrusion Detection System", 36, True, WHITE)

    sub = s.shapes.add_textbox(Inches(1.0), Inches(3.7), Inches(10), Inches(0.6))
    p = sub.text_frame.paragraphs[0]
    run = p.add_run()
    set_run(run, "How NIDS protects your network — architecture, detection, and value", 18, False, MUTED)

    meta = s.shapes.add_textbox(Inches(1.0), Inches(5.8), Inches(10), Inches(0.5))
    p = meta.text_frame.paragraphs[0]
    run = p.add_run()
    set_run(run, "Client Presentation  ·  Confidential", 14, False, MUTED)

    # ----- 2. Agenda -----
    s = blank_slide(prs)
    add_bg(s)
    add_title_block(s, "Agenda", "What we will cover today")
    items = [
        ("01", "What is NIDS?"),
        ("02", "Why it matters for your business"),
        ("03", "Architecture & data flow"),
        ("04", "Core components"),
        ("05", "Detection methods"),
        ("06", "NIDS vs HIDS & IPS"),
        ("07", "Deployment models"),
        ("08", "Business benefits & next steps"),
    ]
    for i, (num, label) in enumerate(items):
        col = i % 2
        row = i // 2
        left = Inches(0.9 + col * 6.0)
        top = Inches(1.7 + row * 1.15)
        num_box = s.shapes.add_textbox(left, top, Inches(0.8), Inches(0.5))
        p = num_box.text_frame.paragraphs[0]
        run = p.add_run()
        set_run(run, num, 22, True, TEAL)
        lab = s.shapes.add_textbox(left + Inches(0.9), top + Inches(0.05), Inches(4.5), Inches(0.5))
        p = lab.text_frame.paragraphs[0]
        run = p.add_run()
        set_run(run, label, 18, False, NAVY)
    add_footer(s, 2, total)

    # ----- 3. What is NIDS -----
    s = blank_slide(prs)
    add_bg(s)
    add_title_block(s, "What is a Network IDS?", "Passive monitoring that finds threats on the wire")
    add_card(
        s,
        Inches(0.7),
        Inches(1.6),
        Inches(5.8),
        Inches(4.8),
        "In plain terms",
        [
            "A Network Intrusion Detection System (NIDS) watches network traffic in real time.",
            "",
            "It sits outside the main path (usually on a tap or SPAN port) so it does not slow production traffic.",
            "",
            "When suspicious patterns appear — known attacks, scans, malware callbacks — it raises alerts for your security team.",
        ],
    )
    add_card(
        s,
        Inches(6.8),
        Inches(1.6),
        Inches(5.8),
        Inches(4.8),
        "What it is not",
        [
            "Not a firewall — it does not block by default.",
            "",
            "Not antivirus on endpoints — it watches network flows, not files on laptops.",
            "",
            "Not a replacement for patching or access control — it complements them as an early-warning layer.",
        ],
        icon_color=TEAL_DARK,
    )
    add_footer(s, 3, total)

    # ----- 4. Why it matters -----
    s = blank_slide(prs)
    add_bg(s)
    add_title_block(s, "Why NIDS matters", "Visibility into threats that slip past perimeter controls")
    cards = [
        ("Early warning", "Detect reconnaissance, lateral movement, and C2 traffic before damage spreads."),
        ("Blind-spot coverage", "Catch threats that bypassed firewalls, VPN, or compromised credentials."),
        ("Compliance support", "Evidence of continuous monitoring for audits (ISO, SOC 2, PCI, etc.)."),
        ("Incident response", "Rich packet/flow context helps analysts investigate faster."),
    ]
    for i, (t, b) in enumerate(cards):
        col = i % 2
        row = i // 2
        add_card(
            s,
            Inches(0.7 + col * 6.2),
            Inches(1.6 + row * 2.4),
            Inches(5.9),
            Inches(2.2),
            t,
            [b],
        )
    add_footer(s, 4, total)

    # ----- 5. Architecture diagram -----
    s = blank_slide(prs)
    add_bg(s)
    add_title_block(s, "NIDS architecture", "End-to-end traffic path from network to analyst")
    if ARCH_IMG.exists():
        # Fit image under title
        s.shapes.add_picture(str(ARCH_IMG), Inches(0.8), Inches(1.45), width=Inches(11.7))
    else:
        box = s.shapes.add_textbox(Inches(1), Inches(3), Inches(10), Inches(1))
        p = box.text_frame.paragraphs[0]
        run = p.add_run()
        set_run(run, "[Architecture diagram]", 18, False, MUTED)
    add_footer(s, 5, total)

    # ----- 6. How it works -----
    s = blank_slide(prs)
    add_bg(s)
    add_title_block(s, "How it works", "Five steps from packets to action")
    steps = [
        ("1", "Capture", "Traffic is mirrored via SPAN/tap from critical network segments."),
        ("2", "Decode", "Sensor parses protocols (TCP/UDP, HTTP, DNS, TLS metadata, etc.)."),
        ("3", "Detect", "Engine matches signatures and/or scores anomalous behavior."),
        ("4", "Alert", "High-confidence events go to SIEM / SOC console."),
        ("5", "Respond", "Analysts triage, contain, and tune rules for fewer false positives."),
    ]
    for i, (n, t, b) in enumerate(steps):
        left = Inches(0.5 + i * 2.5)
        # circle number
        circ = s.shapes.add_shape(MSO_SHAPE.OVAL, left + Inches(0.75), Inches(1.8), Inches(0.7), Inches(0.7))
        circ.fill.solid()
        circ.fill.fore_color.rgb = TEAL
        circ.line.fill.background()
        nb = s.shapes.add_textbox(left + Inches(0.75), Inches(1.92), Inches(0.7), Inches(0.5))
        p = nb.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        set_run(run, n, 20, True, WHITE)
        if i < 4:
            line = s.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                left + Inches(1.6),
                Inches(2.1),
                Inches(1.5),
                Inches(0.06),
            )
            line.fill.solid()
            line.fill.fore_color.rgb = LIGHT_LINE
            line.line.fill.background()
        tb = s.shapes.add_textbox(left, Inches(2.8), Inches(2.35), Inches(0.4))
        p = tb.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        set_run(run, t, 16, True, NAVY)
        bb = s.shapes.add_textbox(left, Inches(3.3), Inches(2.35), Inches(2.2))
        tf = bb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        set_run(run, b, 12, False, BODY)
    add_footer(s, 6, total)

    # ----- 7. Components -----
    s = blank_slide(prs)
    add_bg(s)
    add_title_block(s, "Core components", "Building blocks of a production NIDS")
    comps = [
        ("Network tap / SPAN", "Copies traffic without interrupting production paths."),
        ("NIDS sensor", "Appliance or VM that receives and preprocesses packets."),
        ("Detection engine", "Signature matching + optional ML/anomaly analytics."),
        ("Rule / intel feeds", "Vendor signatures, threat intel IOCs, custom rules."),
        ("Alert manager", "Deduplicates, prioritizes, and routes events."),
        ("SIEM / console", "Where SOC analysts hunt, triage, and report."),
    ]
    for i, (t, b) in enumerate(comps):
        col = i % 3
        row = i // 3
        add_card(
            s,
            Inches(0.55 + col * 4.15),
            Inches(1.55 + row * 2.5),
            Inches(4.0),
            Inches(2.3),
            t,
            [b],
        )
    add_footer(s, 7, total)

    # ----- 8. Detection methods -----
    s = blank_slide(prs)
    add_bg(s)
    add_title_block(s, "Detection methods", "Two complementary approaches")
    add_card(
        s,
        Inches(0.7),
        Inches(1.6),
        Inches(5.8),
        Inches(4.8),
        "Signature-based",
        [
            "Matches traffic against known attack patterns (CVE exploits, malware beacons, scanners).",
            "",
            "Strengths: high precision for known threats; easy to explain.",
            "",
            "Limits: weak against brand-new (zero-day) attacks until rules update.",
        ],
    )
    add_card(
        s,
        Inches(6.8),
        Inches(1.6),
        Inches(5.8),
        Inches(4.8),
        "Anomaly / behavior-based",
        [
            "Baselines normal traffic, then flags unusual volumes, destinations, or protocols.",
            "",
            "Strengths: can surface unknown or insider threats.",
            "",
            "Limits: needs tuning; more false positives without good baselines.",
        ],
        icon_color=TEAL_DARK,
    )
    add_footer(s, 8, total)

    # ----- 9. NIDS vs others -----
    s = blank_slide(prs)
    add_bg(s)
    add_title_block(s, "NIDS vs HIDS vs IPS", "Choosing the right control for the job")

    # simple table via cards
    headers = ["", "NIDS", "HIDS", "IPS"]
    rows = [
        ["Watches", "Network traffic", "Host/OS activity", "Network (inline)"],
        ["Placement", "Tap / SPAN", "On each server", "In traffic path"],
        ["Primary action", "Detect & alert", "Detect on host", "Detect & block"],
        ["Best for", "East-west & perimeter visibility", "Rootkits, file integrity", "Automated prevention"],
    ]
    # header row
    col_w = [Inches(2.4), Inches(3.2), Inches(3.2), Inches(3.2)]
    lefts = [Inches(0.6)]
    for w in col_w[:-1]:
        lefts.append(lefts[-1] + w)

    for c, h in enumerate(headers):
        shape = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, lefts[c], Inches(1.6), col_w[c] - Inches(0.08), Inches(0.55))
        shape.fill.solid()
        shape.fill.fore_color.rgb = NAVY if c > 0 else OFF_WHITE
        shape.line.fill.background()
        tb = s.shapes.add_textbox(lefts[c] + Inches(0.1), Inches(1.7), col_w[c] - Inches(0.2), Inches(0.4))
        p = tb.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER if c > 0 else PP_ALIGN.LEFT
        run = p.add_run()
        set_run(run, h, 14, True, WHITE if c > 0 else MUTED)

    for r, row in enumerate(rows):
        top = Inches(2.25 + r * 0.95)
        for c, cell in enumerate(row):
            bgc = WHITE if r % 2 == 0 else CARD_BG
            if c == 0:
                bgc = CARD_BG
            shape = s.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, lefts[c], top, col_w[c] - Inches(0.08), Inches(0.88)
            )
            shape.fill.solid()
            shape.fill.fore_color.rgb = bgc
            shape.line.color.rgb = LIGHT_LINE
            shape.line.width = Pt(0.5)
            tb = s.shapes.add_textbox(lefts[c] + Inches(0.12), top + Inches(0.25), col_w[c] - Inches(0.25), Inches(0.5))
            p = tb.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
            run = p.add_run()
            set_run(run, cell, 12, c == 0, NAVY if c == 0 else BODY)

    note = s.shapes.add_textbox(Inches(0.7), Inches(6.2), Inches(11.5), Inches(0.4))
    p = note.text_frame.paragraphs[0]
    run = p.add_run()
    set_run(run, "Recommendation: use NIDS for network visibility; pair with HIDS/EDR on critical hosts and IPS where blocking is required.", 13, False, MUTED)
    add_footer(s, 9, total)

    # ----- 10. Deployment -----
    s = blank_slide(prs)
    add_bg(s)
    add_title_block(s, "Deployment models", "Where to place sensors for maximum coverage")
    models = [
        ("Perimeter / DMZ", "Watch internet-facing traffic for scans, exploits, and malware downloads."),
        ("Data center core", "Monitor east-west traffic between servers — where attackers often move unseen."),
        ("Cloud / hybrid", "Virtual taps or cloud-native flow + packet mirrors for VPC/VNet segments."),
        ("Branch / OT edges", "Lightweight sensors at remote sites or industrial networks when risk warrants."),
    ]
    for i, (t, b) in enumerate(models):
        col = i % 2
        row = i // 2
        add_card(
            s,
            Inches(0.7 + col * 6.2),
            Inches(1.6 + row * 2.4),
            Inches(5.9),
            Inches(2.2),
            t,
            [b],
        )
    add_footer(s, 10, total)

    # ----- 11. Benefits -----
    s = blank_slide(prs)
    add_bg(s)
    add_title_block(s, "Business benefits", "What your organization gains")
    benefits = [
        ("Reduced dwell time", "Spot intrusions earlier and shrink attacker time-to-damage."),
        ("Stronger SOC capability", "Actionable alerts with packet context for faster triage."),
        ("Risk reduction", "Additional control layer when credentials or endpoints fail."),
        ("Audit readiness", "Demonstrable continuous network monitoring."),
        ("Scalable visibility", "One sensor can cover many hosts without agent sprawl."),
        ("Tunable over time", "Rules and baselines improve as your environment is learned."),
    ]
    for i, (t, b) in enumerate(benefits):
        col = i % 3
        row = i // 3
        add_card(
            s,
            Inches(0.55 + col * 4.15),
            Inches(1.55 + row * 2.5),
            Inches(4.0),
            Inches(2.3),
            t,
            [b],
        )
    add_footer(s, 11, total)

    # ----- 12. Next steps / close -----
    s = blank_slide(prs)
    add_bg(s)
    add_title_block(s, "Recommended next steps", "A practical path from briefing to pilot")
    steps_next = [
        ("1. Scope", "Identify critical segments: internet edge, DC core, key cloud VPCs."),
        ("2. Pilot", "Deploy one sensor on SPAN/tap; integrate alerts into your SIEM."),
        ("3. Tune", "Reduce noise for 2–4 weeks; prioritize high-fidelity use cases."),
        ("4. Expand", "Roll out to remaining segments; define SLAs for SOC response."),
    ]
    for i, (t, b) in enumerate(steps_next):
        top = Inches(1.55 + i * 1.15)
        num = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top, Inches(11.7), Inches(1.0))
        num.fill.solid()
        num.fill.fore_color.rgb = WHITE
        num.line.color.rgb = LIGHT_LINE
        try:
            num.adjustments[0] = 0.1
        except Exception:
            pass
        stripe = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), top, Inches(0.1), Inches(1.0))
        stripe.fill.solid()
        stripe.fill.fore_color.rgb = TEAL
        stripe.line.fill.background()
        tb = s.shapes.add_textbox(Inches(1.2), top + Inches(0.15), Inches(10.8), Inches(0.35))
        p = tb.text_frame.paragraphs[0]
        run = p.add_run()
        set_run(run, t, 16, True, NAVY)
        bb = s.shapes.add_textbox(Inches(1.2), top + Inches(0.5), Inches(10.8), Inches(0.35))
        p = bb.text_frame.paragraphs[0]
        run = p.add_run()
        set_run(run, b, 13, False, BODY)
    add_footer(s, 12, total)

    # ----- Closing thank you (bonus as slide 12 is next steps; make closing replace? keep 12 as is)
    # Actually I counted 12 including next steps. Good.

    prs.save(str(OUTPUT))
    print(f"Saved: {OUTPUT}")


if __name__ == "__main__":
    build()
